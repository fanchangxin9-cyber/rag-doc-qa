"""
knowledge.py —— 档案室

职责只有三件：
  1. 文件怎么进来（认后缀 → 解析 → 结构感知切块）
  2. 怎么存（写进向量库）
  3. 怎么查、怎么删

对外只露四个函数：load_db / add_file / list_files / delete_file。

支持的格式：PDF / Word(.docx) / Markdown / TXT。
每种格式写一个 _sections_xxx 解析器，统一吐 [(正文, 位置说词)]；
再交给 _to_documents 收口成入库用的块。
往后想加格式（HTML、Excel…）：加一个解析器 + 在 _PARSERS 里登记一行即可，
下面的入库、查重、删除代码一行都不用动。

注意：这个文件刻意不 import fastapi。
档案室不该知道"网页"这个东西存在 —— 这样它将来也能被命令行工具、
测试脚本直接调用。出错时它抛 KnowledgeError，由 main.py 负责翻译成
给浏览器看的错误。
"""

import os
import re
import shutil

from langchain_community.document_loaders import PyPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_text_splitters import MarkdownHeaderTextSplitter
from langchain_core.documents import Document
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import FastEmbedEmbeddings

import config


class KnowledgeError(Exception):
    """知识库相关的「能被预料到」的错误。

    比如：文件读不出文字、格式不支持、重复上传、要删的文件不存在。
    能被预料的错误，用这个异常明确说出来；
    真正没预料到的 bug（代码写错了）就让它原样炸掉，别在这里吞掉。
    """


# 向量模型：档案室自己的工具，负责"把文字变成一串数字"。
# 每个模块自己养自己需要的模型，互不干扰。
EMBEDDING = FastEmbedEmbeddings(model_name=config.EMBED_MODEL, cache_dir=config.CACHE_DIR)


# ============================================================
# 一、切「节」：每种格式一个解析器，统一吐 [(正文, 位置说词)]
# ============================================================
# 「节」= 一段语义完整、自带位置说明的文字。它还不是最终入库的块
# （可能太长），要在 _to_documents 里再收口一次。
#
# 为什么四个解析器都统一成 (正文, 位置) 这个形状？
# 为了让后面「收口」和「入库」的代码【完全不关心文件是什么格式】。
# 这就是"接口"的价值：约定好形状，零件随便换。

def _read_text(path: str) -> str:
    """读纯文本文件（TXT / Markdown 都走它）。

    中文 TXT 常见两种编码：UTF-8 和 GBK（Windows 记事本的老习惯）。
    先按 UTF-8 试，读失败（抛 UnicodeDecodeError）就换 GBK 再试。
    utf-8-sig 能顺带干掉文件开头那个看不见的 BOM 字符。
    """
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "gbk"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise KnowledgeError("这个文本文件读不出内容：编码既不是 UTF-8 也不是 GBK，请另存为 UTF-8 再传")


def _sections_pdf(path: str) -> list:
    """PDF：天然的分节单位就是「页」，位置说词就是「第 N 页」。"""
    docs = PyPDFLoader(path).load()
    return [(d.page_content, f"第 {d.metadata.get('page', 0) + 1} 页") for d in docs]


def _sections_markdown(path: str) -> list:
    """Markdown：按 # / ## / ### 切，位置说词是「一级 > 二级 > 三级」。

    strip_headers=False 很重要 —— 它让标题文字【留在正文里】。
    否则切完只剩内容、标题被抽走，BM25 就再也搜不到「维护保养」
    这种只出现在标题里的词了。
    """
    splitter = MarkdownHeaderTextSplitter(
        headers_to_split_on=config.MD_HEADERS,
        strip_headers=False,
    )
    out = []
    for d in splitter.split_text(_read_text(path)):
        # metadata 里就是 {'一级': '一、概述', '二级': '1.2 安装'} 这样的键值对
        path_txt = " > ".join(d.metadata[k] for _, k in config.MD_HEADERS if k in d.metadata)
        out.append((d.page_content, path_txt))     # 没有标题的文档，路径就是空串
    return out


def _para_label(a: int, b: int) -> str:
    """段号说词：第 3 段 / 第 3–7 段"""
    return f"第 {a} 段" if a == b else f"第 {a}–{b} 段"


def _sections_txt(path: str) -> list:
    """TXT：没有标题可依，退一步用「空行」认段落。

    做法是「攒包」：从头把段落一个个装进书包，装不下（超过 CHUNK_SIZE）就先封包。
    这样每个包都是完整段落拼成的，位置也能说清是「第 3–7 段」。
    """
    paras = [p.strip() for p in re.split(r"\n\s*\n", _read_text(path)) if p.strip()]

    # start 记的是"当前这一包的起始段号"，段号从 1 开始数，所以初值是 1。
    packs, buf, start, size = [], [], 1, 0
    for i, p in enumerate(paras, start=1):
        if buf and size + len(p) > config.CHUNK_SIZE:      # 装不下了 → 先封这一包
            packs.append(("\n\n".join(buf), _para_label(start, i - 1)))
            buf, size, start = [], 0, i                    # 清空书包，下一包从第 i 段起
        buf.append(p)
        size += len(p)
    if buf:                                                # 最后一包
        packs.append(("\n\n".join(buf), _para_label(start, len(paras))))
    return packs


_HEADING_RE = re.compile(r"(?:Heading|标题)\s*(\d+)", re.IGNORECASE)


def _heading_level(style_name: str) -> int:
    """把 Word 的样式名翻译成标题层级。

    英文版 Word 叫「Heading 1」，中文版模板叫「标题 1」，两种都认。
    不是标题样式就返回 0（表示"这是普通正文"）。
    """
    m = _HEADING_RE.search(style_name or "")
    return int(m.group(1)) if m else 0


def _path_label(stack: dict) -> str:
    """把标题栈拼成「一章 > 1.2 节」这样的人话。

    一个标题都没有（文档开头那段正文）就返回空串 ——
    空串会被理解成「这段文字没有章节可归属」，界面上只显示文件名。
    """
    return " > ".join(stack[k] for k in sorted(stack))


# 两条「这看起来像标题」的线索。正则里的 ^\s* 是容忍标题前有空格。
#   装饰符号：▎ ■ ◆ ▲ ★ ▶ ► § ¶ 【 …  —— 特意【不】收 ● • - * 这几个，
#   因为它们太常被当项目符号用了，一收就会把要点列表误认成一堆标题。
_MARKER_RE = re.compile(r"^\s*[▎■◆▲★▶►§¶]|^\s*[【〔〖]")
#   章节编号：第 3 章 / 第二节 / 一、 / 1.2
#   数字限定 1~2 位：不这么写的话，「2023.09——2027.06」这类日期也会被当成编号。
#   【只认多级编号（1.2 / 1.2.3），不认单独的「1.」「1、」】——
#   后者在中文文档里当【列表】用的频率远高于当章节名（「1. 关闭电源 / 2. 拔下插头」），
#   认了会把一份操作步骤切成十几个假章节，还会把内容切丢。宁可漏认，也不能误认。
_NUMBER_RE = re.compile(
    r"^\s*(?:第\s*[一二三四五六七八九十百\d]+\s*[章节篇部]"
    r"|[一二三四五六七八九十]+\s*[、．.]"
    r"|\d{1,2}(?:\.\d{1,2})+)"
)


def _looks_like_manual_heading(text: str) -> bool:
    """猜「这是不是一段手工排版的标题」（没套标题样式、靠排版凑出来的）。

    三个条件必须【同时】满足，缺一个都不算：
      ① 短：不超过 config.DOCX_HEADING_MAX_CHARS 字；
      ② 不以句末标点结尾 —— 标题几乎不用「。」收尾，而正文句子几乎一定用；
      ③ 开头带「标题味」：装饰符号，或章节编号。

    为什么定得这么保守？因为【宁可漏认，也不能误认】：
    漏认顶多退回按字数切；误认会把一段完整内容切得七零八落，
    还凭空造出一堆假章节名，比不认更糟。
    所以不带标志、只是"短且加粗"的段落，这里【不】认 —— 那种在简历里
    多半是「技能：xxx」这类小标签，认了反而坏事。
    """
    t = text.strip()
    if not t or len(t) > config.DOCX_HEADING_MAX_CHARS:
        return False
    if t[-1] in "。！？；，、：!?;,.":
        return False
    return bool(_MARKER_RE.match(t) or _NUMBER_RE.match(t))


def _docx_blocks(doc):
    """按【正文的真实顺序】把 Word 里的「段落」和「表格」逐个吐出来。

    为什么不直接用 doc.paragraphs？
      ① 那个列表里【完全没有表格】，整张表会被丢掉；
      ② 段落和表格混排时，顺序也不保证和正文一致。
    遍历底层 XML 的 body 子节点，才能保证顺序不乱。
    表格被转成「a | b」每行一条的纯文本，当成普通段落参与切块。
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)          # 段落对象
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
            yield "\n".join(rows)                # 表格 → 纯文本字符串


def _sections_docx(path: str) -> list:
    """Word：靠「标题」认章节，维护一条「当前章节路径」。

    标题有两个来源：
      ① 正经的标题样式 ——「标题 1」「Heading 2」这类，看样式的名字；
      ② 手工排版的标题 —— 没套样式，只是打了个 ▎ 或者写成「一、」，
         由 _looks_like_manual_heading() 猜出来（可用 config.DOCX_GUESS_HEADINGS 关掉）。
    现实里的中文文档，②占了很大一部分，不认它"结构感知"就基本等于没做。

    思路很像逛一栋楼：
      · 遇到一级标题 → 换到新的一层（更深的旧标题全部作废）；
      · 遇到二级标题 → 在当前层里换房间；
      · 遇到正文     → 装进当前房间的筐里；
      · 再次遇到标题 → 先把筐里的东西交出去（封一个节），再换房间。

    这里把 python-docx 的导入放在函数【里面】，而不是文件顶部：
    这样就算没装 python-docx，PDF / Markdown / TXT 照样能用，
    只有传 Word 时才给一句明确的提示，而不是整个服务起不来。
    """
    try:
        from docx import Document as WordDocument
    except ImportError:
        raise KnowledgeError("环境里没装 python-docx，读不了 Word 文件；请先执行 pip install python-docx")

    doc = WordDocument(path)
    sections, buf, stack = [], [], {}

    def flush():
        """把筐里的正文封成一个节。

        兜底（很重要）：筐虽然是空的，但只要已经进过章节（stack 非空），
        就说明「上一行标题自己的文字」还挂在章节路径里、从没被交出去过 ——
        这时必须把它单独封成一个只有标题的小节，否则【这一行字会静默消失】。
        （手工编号列表最容易踩：每行都被当成标题、行行后面都没有正文。）

        一句话：宁可多切出一个"只有标题"的小节，也绝不能凭空少掉一行字。
        """
        path = _path_label(stack)
        body = "\n".join(buf)
        if not path and not body:
            return                                 # 没进过章节、筐也空 —— 真没什么可交
        # 关键：把「章节路径」这一行【写进正文里】，跟 Markdown 的做法保持一致。
        # 只把路径放进元数据是不够的 —— BM25 是拿正文的文字建词表的，
        # 标题不落在正文里，「总则」「维护保养」这类只出现在标题里的词就永远搜不到。
        sections.append((f"{path}\n{body}" if path and body else (path or body), path))
        buf.clear()                                # 书包清空，下一节重新装

    for block in _docx_blocks(doc):
        if isinstance(block, str):                 # 表格转出来的纯文本
            buf.append(block)
            continue

        level = _heading_level(block.style.name)          # ① 正经的标题样式（标题 1 / Heading 2）
        if not level and config.DOCX_GUESS_HEADINGS and _looks_like_manual_heading(block.text):
            level = 1                                     # ② 猜出来的手工标题，一律当一级
        if level:
            flush()
            stack[level] = block.text.strip()
            for deeper in [k for k in stack if k > level]:   # 换了新章节，更深的旧标题作废
                del stack[deeper]
        elif block.text.strip():
            buf.append(block.text.strip())

    flush()
    return sections


# 分派表：后缀 → 解析器。想支持新格式，写一个 _sections_xxx 再加一行即可。
_PARSERS = {
    ".pdf": _sections_pdf,
    ".docx": _sections_docx,
    ".md": _sections_markdown,
    ".markdown": _sections_markdown,
    ".txt": _sections_txt,
}


def _to_documents(sections: list, filename: str) -> list:
    """把 [(正文, 位置)] 收口成能入库的 Document 列表 —— 全项目唯一的切块出口。

    规则：
      · 一节不超过 CHUNK_SIZE 字 → 原样就是一块（这就是「结构感知」的收益：
        章节短就整章一块，不会被字数刀切成两半）；
      · 一节太长（比如一整章几千字）→ 才按字数再切，切出的每一块
        都【继承同一节的位置】—— 位置说的是"这段话出自哪里"，
        同一节里切出来的碎块，出处当然一样。
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
    )
    out = []
    for text, location in sections:
        text = text.strip()
        if not text:
            continue                                   # 空节直接跳过
        pieces = [text] if len(text) <= config.CHUNK_SIZE else splitter.split_text(text)
        for piece in pieces:
            out.append(Document(
                page_content=piece,
                # 元数据只留两样，而且【所有格式长得一模一样】
                metadata={"file_name": filename, "location": location},
            ))
    return out


# ============================================================
# 二、入库 / 查库 / 删库
# ============================================================
def load_db():
    """取出向量库。库里还是空的，就返回 None（由调用方决定怎么提示）"""
    if not os.path.exists(config.DB_PATH):
        return None
    return FAISS.load_local(
        config.DB_PATH,
        EMBEDDING,
        # FAISS 用 pickle 存盘，加载时必须显式同意，否则会报错。
        # 因为库文件是我们自己生成的、不来自外部，这里同意是安全的。
        allow_dangerous_deserialization=True,
    )


def all_chunks(db):
    """把库里所有片段取出来。

    为什么要这么绕？因为 FAISS 没有"列出全部文档"的公开接口，
    只能从内部的 docstore 字典里捞（`_dict` 前面的下划线就表示"内部使用"）。
    整个项目只有这一处碰内部结构 —— 将来换成别的向量库，也只需要改这一个函数。

    对 main.py 公开（不带下划线），是因为关键词检索需要拿到全库片段来建词表。
    """
    return list(db.docstore._dict.values())


def add_file(content: bytes, filename: str) -> dict:
    """把一份文档收进库：认后缀 → 解析切节 → 收口成块 → 贴标签 → 存起来。

    content 是文件的二进制内容，filename 是原始文件名。
    """
    # ① 认后缀：不认识的格式，在这里就拦住，给出"支持哪些"的清单
    ext = os.path.splitext(filename)[1].lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise KnowledgeError(
            f"不支持「{ext or '无扩展名'}」这种文件；目前支持：{'、'.join(config.SUPPORTED_EXTS)}"
        )

    # 先落到一个临时文件上，因为几个解析器都需要"硬盘上的路径"
    temp_path = f"./temp_{filename}"
    with open(temp_path, "wb") as f:
        f.write(content)

    try:
        # ② 解析 + 切节：每种格式各有各的规矩，但吐出来的形状统一
        sections = parser(temp_path)

        # ③ 收口：太长的节再按字数切，并补上 file_name / location 两个标签
        split_docs = _to_documents(sections, filename)

        # 守卫：这份文件一个字都没读出来
        if not split_docs:
            hint = "（可能是扫描件或图片版，请换一份能选中文字的 PDF）" if ext == ".pdf" else ""
            raise KnowledgeError(f"「{filename}」里读不到文字{hint}")

        # ④ 入库：库已存在 → 先查重再追加；不存在 → 新建一个
        db = load_db()
        if db is None:
            db = FAISS.from_documents(split_docs, EMBEDDING)
        else:
            if _already_exists(db, filename):
                raise KnowledgeError(f"「{filename}」已经在知识库里了，不用重复上传")
            db.add_documents(split_docs)

        db.save_local(config.DB_PATH)
    finally:
        # finally = 不管上面成功还是中途报错，这里都一定会执行
        # 所以要删的临时文件放在这里，绝不会漏掉
        if os.path.exists(temp_path):
            os.remove(temp_path)

    return {
        "msg": "文档入库成功",
        "file": filename,
        "chunk_num": len(split_docs),
        "total": db.index.ntotal,
    }


def _already_exists(db, filename: str) -> bool:
    """这个文件名是不是已经入库了（防止重复上传）"""
    for d in all_chunks(db):
        if d.metadata.get("file_name") == filename:
            return True
    return False


def list_files() -> dict:
    """列出库里已经有哪些文件，各自多少段"""
    db = load_db()
    if db is None:
        return {"files": [], "total": 0}

    # 数一数每个文件各有多少段
    counts = {}
    for d in all_chunks(db):
        name = d.metadata.get("file_name", "未知文件")
        counts[name] = counts.get(name, 0) + 1

    files = [{"file": name, "chunk_num": count} for name, count in counts.items()]
    return {"files": files, "total": db.index.ntotal}


def delete_file(filename: str) -> dict:
    """从库里删掉一整份文件。

    FAISS 不能直接删单条向量，所以做法是：把这文件之外的片段全挑出来，
    删掉整个库，再用挑出来的片段重建。
    """
    db = load_db()
    if db is None:
        raise KnowledgeError("知识库还是空的，没有可删除的文件")

    chunks = all_chunks(db)

    # 把「不属于这个文件」的段挑出来，这些是要保留的
    keep = [d for d in chunks if d.metadata.get("file_name") != filename]

    # 保留的数量和原来一样多 → 说明一段都没删掉 → 这个文件根本不在库里
    if len(keep) == len(chunks):
        raise KnowledgeError(f"知识库里没有「{filename}」这个文件")

    shutil.rmtree(config.DB_PATH)
    if keep:
        FAISS.from_documents(keep, EMBEDDING).save_local(config.DB_PATH)

    return {"msg": "已删除", "file": filename, "left": len(keep)}